from datetime import date
from io import BytesIO

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from st_aggrid import AgGrid, DataReturnMode, GridOptionsBuilder, JsCode

from controllers.productos_obsoletos_controller import (
    ALMACENES_DISPONIBLES,
    obtener_reporte_productos_obsoletos_ctrl,
)

_COLS_FECHA = {"Fecha Última Compra", "Fecha Último Movimiento", "Fecha de Caducidad"}

# muy tenues a propósito — son columnas calculadas, no deben competir
# visualmente con los datos de identificación del lote
_AZUL_ENTRADAS = "DCE6F5"
_VERDE_SALIDAS = "E2F0DD"
_MORADO_FINAL = "EAE2F5"

_CELL_STYLE_ENTRADAS_JS = JsCode(f"function(p){{return {{backgroundColor:'#{_AZUL_ENTRADAS}'}};}}")
_CELL_STYLE_SALIDAS_JS = JsCode(f"function(p){{return {{backgroundColor:'#{_VERDE_SALIDAS}'}};}}")
_CELL_STYLE_FINAL_JS = JsCode(f"function(p){{return {{backgroundColor:'#{_MORADO_FINAL}'}};}}")

# Kg: separador de miles, sin decimales. USD: separador de miles, 2 decimales.
_VALUE_FORMATTER_KG_JS = JsCode("""
function(params) {
    if (params.value === null || params.value === undefined || params.value === '') return '';
    return Number(params.value).toLocaleString('en-US', {minimumFractionDigits: 0, maximumFractionDigits: 0});
}
""")
_VALUE_FORMATTER_USD_JS = JsCode("""
function(params) {
    if (params.value === null || params.value === undefined || params.value === '') return '';
    return '$' + Number(params.value).toLocaleString('en-US', {minimumFractionDigits: 2, maximumFractionDigits: 2});
}
""")

_COLS_USD_FIJAS = {"Costo Promedio del Artículo", "Costo Promedio Total"}

# mismo texto para el expander de la app y la hoja "Notas" del Excel — una
# sola fuente, para que no se desalineen con el tiempo
_NOTAS_CALCULO = [
    (
        "Entradas <mes>",
        "Suma de todos los movimientos con signo positivo de ESE lote ocurridos en ese mes "
        "(compras, ajustes de entrada, devoluciones, producción, etc.) — no se limita a compras "
        "facturadas, cuenta cualquier entrada real al lote.",
    ),
    (
        "Salidas <mes>",
        "Suma de todos los movimientos con signo negativo de ESE lote ocurridos en ese mes "
        "(consumo, ventas, mermas, ajustes de salida, etc.), reportada en positivo (unidades "
        "que salieron).",
    ),
    (
        "Final <mes>",
        "Existencia real de HOY menos todo lo que se movió (con su signo) después del cierre "
        "de ese mes — no se asume un saldo inicial, así que siempre cuadra exactamente con la "
        "existencia real actual del lote.",
    ),
    (
        "Fecha Última Compra / Año de compra",
        "Se basan únicamente en movimientos de compra facturada (CVE_CPTO=1 en SAE) — no en "
        "órdenes/requisiciones, que no afectan inventario.",
    ),
    (
        "Meses de Vigencia",
        "Meses completos entre hoy y la fecha de caducidad (LTPD01.FCHCADUC, tal cual está "
        "capturada en SAE, sin calcularla) — negativo si el lote ya caducó.",
    ),
    (
        "Costo Promedio del Artículo",
        "Costo promedio actual del artículo en SAE (INVE01.COSTO_PROM) — el mismo valor para "
        "todos los lotes de ese producto, no es específico del lote.",
    ),
    (
        "Costo Promedio Total",
        "Costo Promedio del Artículo × existencia actual del lote — el valor monetario de lo "
        "que sigue sin moverse en ese lote.",
    ),
    (
        "Lotes incluidos",
        "Solo lotes con existencia real mayor a 0.01 al día de hoy (se descartan residuos de "
        "redondeo de punto flotante que en SAE aparecen como existencia ~0 pero no son cero "
        "exacto).",
    ),
]


def _reporte_a_excel_bytes(df: pd.DataFrame, anio: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Obsoletos {anio}"[:31]

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    columnas = list(df.columns)
    ws.append(columnas)

    fill_por_columna = {}
    for col in columnas:
        if col.startswith("Entradas "):
            fill_por_columna[col] = PatternFill(start_color=_AZUL_ENTRADAS, end_color=_AZUL_ENTRADAS, fill_type="solid")
        elif col.startswith("Salidas "):
            fill_por_columna[col] = PatternFill(start_color=_VERDE_SALIDAS, end_color=_VERDE_SALIDAS, fill_type="solid")
        elif col.startswith("Final "):
            fill_por_columna[col] = PatternFill(start_color=_MORADO_FINAL, end_color=_MORADO_FINAL, fill_type="solid")

    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = header_align
        col_name = columnas[cell.column - 1]
        cell.fill = fill_por_columna.get(col_name, header_fill)
    ws.freeze_panes = "L2"

    for _, fila in df.iterrows():
        valores = []
        for col in columnas:
            v = fila[col]
            if pd.isna(v):
                valores.append(None)
            elif col in _COLS_FECHA and hasattr(v, "date"):
                valores.append(v.date())
            else:
                valores.append(v)
        ws.append(valores)

    fila_final = ws.max_row
    for idx, col_name in enumerate(columnas, start=1):
        fill = fill_por_columna.get(col_name)
        if fill:
            for fila_idx in range(2, fila_final + 1):
                ws.cell(row=fila_idx, column=idx).fill = fill

        es_kg = col_name.startswith(("Entradas ", "Salidas ", "Final "))
        es_usd = col_name in _COLS_USD_FIJAS
        if es_kg or es_usd:
            formato = "#,##0" if es_kg else "$#,##0.00"
            for fila_idx in range(2, fila_final + 1):
                ws.cell(row=fila_idx, column=idx).number_format = formato

        largo = len(str(col_name))
        valores_txt = df[col_name].astype(str).tolist()[:200]
        if valores_txt:
            largo = max(largo, max(len(v) for v in valores_txt))
        ws.column_dimensions[get_column_letter(idx)].width = min(largo + 2, 32)

    ws_notas = wb.create_sheet("Notas")
    ws_notas.append(["columna / concepto", "cómo se calcula"])
    for cell in ws_notas[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
    for titulo, texto in _NOTAS_CALCULO:
        ws_notas.append([titulo, texto])
    for fila_idx in range(2, ws_notas.max_row + 1):
        ws_notas.cell(row=fila_idx, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws_notas.column_dimensions["A"].width = 32
    ws_notas.column_dimensions["B"].width = 100

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _mostrar_grid(df: pd.DataFrame):
    gb = GridOptionsBuilder.from_dataframe(df)
    gb.configure_grid_options(autoSizeStrategy=None)
    gb.configure_default_column(editable=False, resizable=True, filter=True, sortable=True, wrapHeaderText=True)
    gb.configure_column("Almacén", pinned="left", width=90)
    gb.configure_column("Clave", pinned="left", width=100)
    gb.configure_column("Producto", pinned="left", width=200)
    gb.configure_column("Lote", pinned="left", width=130)

    for col in df.columns:
        if col.startswith("Entradas "):
            gb.configure_column(col, cellStyle=_CELL_STYLE_ENTRADAS_JS, valueFormatter=_VALUE_FORMATTER_KG_JS, width=130)
        elif col.startswith("Salidas "):
            gb.configure_column(col, cellStyle=_CELL_STYLE_SALIDAS_JS, valueFormatter=_VALUE_FORMATTER_KG_JS, width=130)
        elif col.startswith("Final "):
            gb.configure_column(col, cellStyle=_CELL_STYLE_FINAL_JS, valueFormatter=_VALUE_FORMATTER_KG_JS, width=130)
        elif col in _COLS_USD_FIJAS:
            gb.configure_column(col, valueFormatter=_VALUE_FORMATTER_USD_JS, width=140)

    AgGrid(
        df,
        gridOptions=gb.build(),
        data_return_mode=DataReturnMode.AS_INPUT,
        fit_columns_on_grid_load=False,
        allow_unsafe_jscode=True,
        height=min(56 + len(df) * 35, 520),
        key="pobs_grid",
    )


def mostrar_tab_productos_obsoletos():
    st.subheader("📦 productos obsoletos — reporte por producto y lote")
    st.caption(
        "Existencia real por lote (se excluyen lotes ya en cero) en Materias Primas, "
        "Producto Terminado y Muestras — con entradas/salidas/saldo mes a mes del año "
        "elegido, para detectar qué lleva tiempo sin moverse."
    )

    with st.expander("¿cómo se calcula cada columna? (esto también va en la hoja \"Notas\" del Excel)"):
        st.markdown("\n".join(f"- **{titulo}**: {texto}" for titulo, texto in _NOTAS_CALCULO))

    c1, c2 = st.columns([1, 2])
    with c1:
        anio_sel = st.number_input(
            "año", min_value=2016, max_value=date.today().year, value=date.today().year,
            step=1, key="pobs_anio",
        )
    with c2:
        almacenes_sel = st.multiselect(
            "almacenes",
            options=list(ALMACENES_DISPONIBLES.keys()),
            default=list(ALMACENES_DISPONIBLES.keys()),
            format_func=lambda c: f"{c} - {ALMACENES_DISPONIBLES[c]}",
            key="pobs_almacenes",
        )

    if st.button("🔎 generar reporte", type="primary", key="pobs_btn_generar"):
        with st.spinner("consultando movimientos en SAE…"):
            df = obtener_reporte_productos_obsoletos_ctrl(int(anio_sel), almacenes_sel)
        st.session_state["pobs_df"] = df
        st.session_state["pobs_anio_generado"] = int(anio_sel)

    df = st.session_state.get("pobs_df")
    if df is None:
        return

    if df.empty:
        st.info("no hay lotes con existencia real en esos almacenes.")
        return

    st.success(f"{len(df)} lote(s) con existencia real.")
    _mostrar_grid(df)

    anio_generado = st.session_state.get("pobs_anio_generado", int(anio_sel))
    excel_bytes = _reporte_a_excel_bytes(df, anio_generado)
    st.download_button(
        "⬇️ descargar reporte en Excel",
        data=excel_bytes,
        file_name=f"reporte_productos_obsoletos_{anio_generado}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="pobs_btn_descargar",
    )
